"""Chuyển báo cáo Markdown sang Word (.docx) theo định dạng đồ án.

Hỗ trợ:
- Trang A4, Times New Roman, giãn dòng 1.5, căn đều hai bên.
- Trang bìa, heading 1-4, ngắt trang theo chương.
- Mục lục, danh mục hình và danh mục bảng dạng field của Word.
- Caption hình từ dòng **Hình X-Y. ...**; caption bảng được sinh tự động.
- Bảng Markdown, ảnh, danh sách, blockquote, code block, chữ đậm/nghiêng/code.
- Ảnh chưa có được thay bằng vị trí chờ để bổ sung sau.
- Header, footer và số trang; Word tự cập nhật field khi mở file.

Chạy mặc định:
    python generate-docx.py

Chạy với đường dẫn chỉ định:
    python generate-docx.py <input.md> <output.docx>

Sau khi mở file Word, nhấn Ctrl+A rồi F9 nếu mục lục/danh mục chưa cập nhật.
"""
import logging
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from PIL import Image

    HAS_PIL = True
except ImportError:
    HAS_PIL = False

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
MAX_IMG_W_CM = 15.5
MAX_IMG_H_CM = 20.0

logger = logging.getLogger(__name__)

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_IMG_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
_FIGURE_CAPTION_RE = re.compile(r"^\*\*(Hình\s+\d+(?:-\d+)?\.\s+.+)\*\*$", re.IGNORECASE)
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+`)")
_NUMBERED_CHAPTER_RE = re.compile(r"^(\d+)\.\s+")
_AUTO_TEXT_RE = re.compile(r"^\*\(Tạo tự động.*\)\*$", re.IGNORECASE)


def set_run_font(run, name=BODY_FONT, size=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_para_shading(paragraph, color):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    p_pr.append(shd)


def keep_table_row(row, repeat_header=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def add_field(paragraph, instruction, placeholder="Nhấn F9 để cập nhật"):
    """Thêm field Word (TOC, PAGE...)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    return run


def add_inline(paragraph, text):
    text = text.replace("&nbsp;", " ")
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, MONO_FONT)
        else:
            run = paragraph.add_run(part)
        if not (part.startswith("`") and part.endswith("`")):
            set_run_font(run)


def img_size_cm(path):
    """Tính kích thước ảnh vừa trang, giữ đúng tỉ lệ."""
    if not HAS_PIL:
        return MAX_IMG_W_CM, None
    with Image.open(path) as image:
        width_px, height_px = image.size
    width_cm = width_px / 96 * 2.54
    height_cm = height_px / 96 * 2.54
    if not width_cm or not height_cm:
        return MAX_IMG_W_CM, None
    scale = min(MAX_IMG_W_CM / width_cm, MAX_IMG_H_CM / height_cm)
    return width_cm * scale, height_cm * scale


def split_table_row(line):
    line = line.strip().removeprefix("|").removesuffix("|")
    return [cell.strip() for cell in line.split("|")]


def read_report_title(md_path):
    """Lấy tên dự án từ heading cấp 2 đầu tiên; fallback về tên file."""
    for line in md_path.read_text(encoding="utf-8").splitlines()[:30]:
        if line.startswith("## "):
            return line[3:].strip()
    return md_path.stem.replace("-", " ").title()


class MdToDocx:
    def __init__(self, md_path: Path):
        self.md_path = md_path
        self.base = md_path.parent
        self.doc = Document()
        self.cover_mode = True
        self.cover_heading_count = 0
        self.current_heading = ""
        self.chapter = None
        self.table_counts = {}
        self._setup_document()

    def _setup_document(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(1.2)
        section.footer_distance = Cm(1.2)
        section.different_first_page_header_footer = True

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = BODY_FONT
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        normal.font.size = Pt(12)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.first_line_indent = Cm(1)
        normal.paragraph_format.widow_control = True

        heading_specs = {
            "Heading 1": (16, True),
            "Heading 2": (14, True),
            "Heading 3": (13, True),
            "Heading 4": (12, True),
        }
        for style_name, (size, bold) in heading_specs.items():
            style = styles[style_name]
            style.font.name = BODY_FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.widow_control = True

        for name in ("Caption Hình", "Caption Bảng"):
            if name not in styles:
                styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style = styles[name]
            style.font.name = BODY_FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            style.font.size = Pt(11)
            style.font.italic = True
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_after = Pt(4)

        self.doc.core_properties.title = "Báo cáo đồ án Kỹ thuật phần mềm"
        self.doc.core_properties.subject = "Thiết kế hệ thống RAG tra cứu Luật Giao thông"
        self.doc.core_properties.author = "Nhóm thực hiện đồ án"

        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        self.doc.settings.element.append(update_fields)

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("Hệ thống RAG tra cứu Luật Giao thông")
        set_run_font(run, size=9)
        run.font.color.rgb = RGBColor(100, 100, 100)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Trang ")
        set_run_font(run, size=10)
        page_run = add_field(footer, " PAGE ", "1")
        set_run_font(page_run, size=10)

    def add_heading(self, text, markdown_level):
        if self.cover_mode and text == "MỤC LỤC":
            self.cover_mode = False

        if self.cover_mode:
            self._add_cover_heading(text, markdown_level)
            return

        word_level = min(markdown_level, 4)
        paragraph = self.doc.add_heading(text, level=word_level)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if word_level == 1:
            paragraph.paragraph_format.page_break_before = True
            chapter_match = _NUMBERED_CHAPTER_RE.match(text)
            self.chapter = chapter_match.group(1) if chapter_match else None
            if self.chapter:
                self.table_counts.setdefault(self.chapter, 0)
        self.current_heading = text

        if text == "MỤC LỤC":
            self._add_index_field('TOC \\o "1-3" \\h \\z \\u')
        elif text == "DANH MỤC HÌNH":
            self._add_index_field('TOC \\h \\z \\t "Caption Hình,1"')
        elif text == "DANH MỤC BẢNG":
            self._add_index_field('TOC \\h \\z \\t "Caption Bảng,1"')

    def _add_cover_heading(self, text, markdown_level):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(18 if self.cover_heading_count else 70)
        paragraph.paragraph_format.space_after = Pt(18)
        run = paragraph.add_run(text)
        set_run_font(run, size=18 if markdown_level == 1 else 16)
        run.bold = True
        self.cover_heading_count += 1

    def _add_index_field(self, instruction):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        add_field(paragraph, instruction)
        note = self.doc.add_paragraph("Mở file bằng Microsoft Word, nhấn Ctrl+A rồi F9 để cập nhật.")
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.first_line_indent = Cm(0)
        for run in note.runs:
            set_run_font(run, size=9)
            run.italic = True
            run.font.color.rgb = RGBColor(120, 120, 120)

    def add_paragraph(self, text):
        figure_match = _FIGURE_CAPTION_RE.match(text)
        if figure_match:
            paragraph = self.doc.add_paragraph(style="Caption Hình")
            add_inline(paragraph, figure_match.group(1))
            return

        paragraph = self.doc.add_paragraph()
        if self.cover_mode:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
        add_inline(paragraph, text)

    def add_bullet(self, text, level=0):
        paragraph = self.doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.75 + 0.5 * level)
        paragraph.paragraph_format.first_line_indent = Cm(-0.4)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        add_inline(paragraph, text)

    def add_number(self, text):
        paragraph = self.doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.4)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        add_inline(paragraph, text)

    def add_quote(self, text):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(1)
        paragraph.paragraph_format.right_indent = Cm(1)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        set_para_shading(paragraph, "F2F2F2")
        add_inline(paragraph, text)
        for run in paragraph.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(85, 85, 85)

    def add_code_block(self, lines):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.right_indent = Cm(0.5)
        paragraph.paragraph_format.line_spacing = 1.0
        set_para_shading(paragraph, "F2F2F2")
        for index, line in enumerate(lines):
            run = paragraph.add_run(line)
            set_run_font(run, MONO_FONT, 9)
            if index < len(lines) - 1:
                run.add_break()

    def add_image(self, alt, rel_path):
        img_path = (self.base / rel_path).resolve()
        if not img_path.exists():
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(12)
            set_para_shading(paragraph, "F2F2F2")
            label = alt or rel_path
            run = paragraph.add_run(f"[ {label} — ảnh sẽ được bổ sung sau ]")
            set_run_font(run, size=10)
            run.italic = True
            run.font.color.rgb = RGBColor(120, 120, 120)
            return

        width_cm, height_cm = img_size_cm(img_path)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run()
        if height_cm:
            run.add_picture(str(img_path), width=Cm(width_cm), height=Cm(height_cm))
        else:
            run.add_picture(str(img_path), width=Cm(width_cm))

    def add_table(self, header, rows):
        self._add_table_caption(header)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for col_index, heading in enumerate(header):
            cell = table.rows[0].cells[col_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, heading)
            for run in paragraph.runs:
                run.bold = True
                set_run_font(run, size=10)
            set_cell_shading(cell, "D9E2F3")
        keep_table_row(table.rows[0], repeat_header=True)

        for row_index, row in enumerate(rows, 1):
            keep_table_row(table.rows[row_index])
            for col_index in range(len(header)):
                value = row[col_index] if col_index < len(row) else ""
                cell = table.rows[row_index].cells[col_index]
                cell.text = ""
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                add_inline(paragraph, value)
                for run in paragraph.runs:
                    set_run_font(run, size=10)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def _add_table_caption(self, header):
        if not self.chapter:
            return
        self.table_counts[self.chapter] = self.table_counts.get(self.chapter, 0) + 1
        number = self.table_counts[self.chapter]
        title = re.sub(r"^\d+(?:\.\d+)*\s*[-–.]?\s*", "", self.current_heading).strip()
        if not title:
            title = " – ".join(header[:2])
        paragraph = self.doc.add_paragraph(style="Caption Bảng")
        run = paragraph.add_run(f"Bảng {self.chapter}-{number}. {title}")
        set_run_font(run, size=11)


def convert(md_path: Path, out_path: Path):
    if not md_path.exists():
        raise FileNotFoundError(f"Không tìm thấy file Markdown: {md_path}")

    converter = MdToDocx(md_path)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            converter.add_code_block(block)
            index += 1
            continue

        if not stripped or stripped in ("---", "***", "___"):
            index += 1
            continue

        if _AUTO_TEXT_RE.match(stripped):
            index += 1
            continue

        match = _HEADING_RE.match(line)
        if match:
            converter.add_heading(match.group(2).strip(), len(match.group(1)))
            index += 1
            continue

        match = _IMG_RE.match(stripped)
        if match:
            converter.add_image(match.group(1).strip(), match.group(2).strip())
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and _TABLE_SEP_RE.match(lines[index + 1]):
            header = split_table_row(lines[index])
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            converter.add_table(header, rows)
            continue

        if stripped.startswith(">"):
            converter.add_quote(stripped.lstrip(">").strip())
            index += 1
            continue

        match = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if match:
            converter.add_bullet(match.group(2).strip(), len(match.group(1)) // 2)
            index += 1
            continue

        match = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if match:
            converter.add_number(match.group(1).strip())
            index += 1
            continue

        converter.add_paragraph(stripped)
        index += 1

    try:
        converter.doc.save(str(out_path))
    except PermissionError as exc:
        raise PermissionError(
            f"Không thể ghi file '{out_path}'. Hãy đóng file nếu đang mở trong Word rồi chạy lại."
        ) from exc
    logger.info("Đã tạo: %s", out_path)
    logger.info("Mở file trong Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục và danh mục.")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)-7s | %(name)s | %(message)s")
    if len(sys.argv) >= 3:
        source = Path(sys.argv[1])
        destination = Path(sys.argv[2])
    else:
        source = Path(__file__).parent / "tai-lieu-thiet-ke-he-thong.md"
        destination = Path(__file__).parent / "tai-lieu-thiet-ke-he-thong.docx"
    convert(source, destination)
