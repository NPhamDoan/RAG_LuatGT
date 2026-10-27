"""Generate bao-cao-phase2-poc.docx từ nội dung Phase 2 POC (Hybrid Search)."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path

doc = Document()

# --- Page margins ---
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# --- Style setup ---
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)


def set_cell_shading(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = tcPr.makeelement(qn("w:shd"), {
        qn("w:fill"): color, qn("w:val"): "clear", qn("w:color"): "auto",
    })
    tcPr.append(shading)


def add_table(headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        set_cell_shading(cell, "D9E2F3")
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = tbl.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(11)
            run.font.name = "Times New Roman"
    doc.add_paragraph()


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def p(text):
    doc.add_paragraph(text)


def bold_p(text):
    para = doc.add_paragraph()
    para.add_run(text).bold = True


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def mixed(parts):
    para = doc.add_paragraph()
    for text, bold in parts:
        run = para.add_run(text)
        run.bold = bold


# ============================================================
# TITLE
# ============================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("BÁO CÁO ĐIỀU TRA TÍNH KHẢ THI")
run.bold = True
run.font.size = Pt(18)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("PHASE 2: Hybrid Search (BM25 + RRF)")
run.bold = True
run.font.size = Pt(14)

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub2.add_run("Hệ thống RAG Tra cứu Luật Giao Thông Việt Nam")
run.italic = True
run.font.size = Pt(12)

doc.add_paragraph()

# ============================================================
# 1. BÀI TOÁN
# ============================================================
h1("1. Bài toán")

h2("1.1. Vì sao cần Phase 2")
p("Phase 1 tìm kiếm hoàn toàn bằng vector (cosine similarity). Cách này hiểu tốt ngữ "
  "nghĩa (\"chạy nhanh trong phố\" gần với \"vượt quá tốc độ khu dân cư\") nhưng lại yếu "
  "với tìm chính xác theo từ khóa / con số:")
bullet("Hỏi \"Điều 650 thừa kế\" → vector không ưu tiên con số \"650\", dễ trả về các Điều "
       "khác cùng chủ đề thừa kế.")
bullet("Hỏi \"tốc độ 60km/h khu dân cư\" → con số \"60\" bị embedding làm mờ, có thể miss "
       "đúng Điều quy định tốc độ.")
p("Phase 2 bổ sung tìm kiếm từ khóa (BM25) chạy song song với vector, rồi hợp nhất hai "
  "bảng xếp hạng bằng RRF (Reciprocal Rank Fusion). Kết quả: giữ nguyên khả năng hiểu "
  "nghĩa của vector, đồng thời bắt được exact-match mà vector bỏ sót.")

h2("1.2. Mục tiêu")
bullet("Tăng độ chính xác truy xuất — đặc biệt với số Điều, số hiệu văn bản, con số (km/h, kg...).")
bullet("KHÔNG thay đổi giao diện, API, prompt, cấu hình. Chỉ nâng chất lượng tầng tìm kiếm.")
bullet("Giữ nguyên logic gating và độ tin cậy của Phase 1 (ngưỡng 0.3 / 0.5 / 0.8).")

h2("1.3. Ràng buộc")
add_table(
    ["Ràng buộc", "Chi tiết"],
    [
        ["Không đổi hợp đồng API", "Frontend, endpoint, schema phản hồi giữ nguyên"],
        ["Không đổi gating", "Điểm số mỗi kết quả vẫn là cosine similarity để so ngưỡng"],
        ["Không thêm dịch vụ ngoài", "BM25 + RRF viết bằng Python thuần, không thêm thư viện, không gọi API"],
        ["Hiệu năng", "Chỉ mục BM25 dựng một lần, cache lại, dựng lại khi dữ liệu đổi"],
    ],
)

# ============================================================
# 2. KHÁI NIỆM NỀN TẢNG
# ============================================================
h1("2. Khái niệm nền tảng")

h2("2.1. Hạn chế của vector search (nhắc lại)")
p("Vector search chuyển câu hỏi và đoạn luật thành vector rồi so độ tương đồng ngữ nghĩa. "
  "Vấn đề: con số và từ khóa hiếm (\"650\", \"60km/h\", \"38/2024/TT-BGTVT\") bị \"hòa tan\" "
  "vào nghĩa chung của câu, nên không được ưu tiên đúng mức. Đây chính là điểm yếu cần bù.")

h2("2.2. BM25 là gì?")
p("BM25 (Okapi BM25) là thuật toán xếp hạng theo TỪ KHÓA (keyword) kinh điển trong tìm "
  "kiếm văn bản. Ý tưởng: một đoạn được coi là khớp nếu chứa ĐÚNG các từ trong câu hỏi, "
  "với ba yếu tố điều chỉnh:")
bullet("TF (term frequency): từ khóa xuất hiện càng nhiều trong đoạn → điểm càng cao (bão hòa dần).")
bullet("IDF (inverse document frequency): từ càng hiếm trong kho → càng quan trọng (vd \"650\" hiếm hơn \"quy định\").")
bullet("Chuẩn hóa độ dài: đoạn dài không bị lợi thế oan chỉ vì chứa nhiều từ.")
p("Công thức dùng hai tham số chuẩn: k1 = 1.5 (điều chỉnh độ bão hòa TF) và b = 0.75 "
  "(mức chuẩn hóa độ dài). BM25 rất mạnh ở đúng chỗ vector yếu: khớp chính xác từ khóa "
  "và con số.")

h2("2.3. Inverted index (chỉ mục ngược)")
p("Để BM25 chạy nhanh, ta dựng inverted index: bản đồ từ mỗi token → danh sách các đoạn "
  "chứa token đó (kèm số lần xuất hiện). Khi truy vấn, chỉ cần duyệt các đoạn chứa token "
  "trong câu hỏi thay vì quét toàn bộ kho.")

h2("2.4. Tokenize + bỏ dấu")
p("Trước khi lập chỉ mục, văn bản được tokenize:")
bullet("Chuyển thường (lowercase).")
bullet("Bỏ dấu tiếng Việt (\"Điều\" → \"dieu\", \"thừa kế\" → \"thua ke\"), để câu có dấu / "
       "không dấu vẫn khớp cùng token.")
bullet("Giữ nguyên chữ số (\"650\", \"60\") — phần quan trọng nhất giúp bắt exact-match.")
p("Ví dụ: \"Điều 650 thừa kế\" → [\"dieu\", \"650\", \"thua\", \"ke\"].")

h2("2.5. RRF là gì? (Reciprocal Rank Fusion)")
p("Sau khi có HAI bảng xếp hạng (một của vector, một của BM25), cần gộp lại thành một danh "
  "sách duy nhất. RRF làm việc này dựa trên THỨ HẠNG chứ không dựa trên điểm số thô (vì "
  "điểm cosine và điểm BM25 không cùng thang đo).")
p("Với mỗi đoạn, điểm RRF = tổng của 1 / (k + thứ_hạng) trên từng bảng mà nó xuất hiện, "
  "với hằng số k = 60 (giá trị chuẩn theo paper Cormack 2009):")
bullet("Đoạn đứng đầu ở cả hai bảng → điểm RRF cao nhất.")
bullet("Đoạn chỉ BM25 bắt được (vector miss) vẫn được cộng điểm → có cơ hội lọt top.")
bullet("Hằng số k = 60 làm mượt: giảm ảnh hưởng của chênh lệch thứ hạng ở đuôi danh sách.")
p("RRF đơn giản, không cần huấn luyện, và là chuẩn công nghiệp để trộn nhiều nguồn xếp hạng.")

h2("2.6. Vì sao \"hybrid\" tốt hơn từng cái riêng")
add_table(
    ["Truy vấn", "Vector đơn thuần", "BM25 đơn thuần", "Hybrid (RRF)"],
    [
        ["\"chạy nhanh trong phố\" (ngữ nghĩa)", "Tốt", "Kém (khác từ)", "Tốt"],
        ["\"Điều 650\" (exact-match)", "Kém", "Tốt", "Tốt"],
        ["\"tốc độ 60km/h khu dân cư\"", "Trung bình", "Tốt (\"60\")", "Tốt"],
    ],
)
p("Hybrid lấy điểm mạnh của cả hai: vector lo phần nghĩa, BM25 lo phần chữ và số.")

# ============================================================
# 3. CÁCH TIẾP CẬN CHI TIẾT
# ============================================================
h1("3. Cách tiếp cận chi tiết")

h2("3.1. Kiến trúc tầng tìm kiếm sau Phase 2")
p("Toàn bộ thay đổi gói gọn trong tầng lưu trữ (vector_store.py). Khi query pipeline gọi "
  "tìm kiếm và CÓ kèm nội dung câu hỏi (query_text), hệ thống chạy nhánh hybrid; nếu không, "
  "vẫn chạy vector-only như Phase 1.")

h2("3.2. Quy trình Hybrid Search")
add_table(
    ["Bước", "Tên", "Xử lý cụ thể"],
    [
        ["1", "Lấy ứng viên vector", "Truy vấn ChromaDB lấy pool rộng max(k×4, 40) đoạn gần nhất kèm điểm cosine"],
        ["2", "Lấy ứng viên BM25", "Tokenize câu hỏi → tra inverted index → xếp hạng BM25, lấy cùng số ứng viên"],
        ["3", "Hợp nhất bằng RRF", "Với mỗi đoạn, cộng 1/(60 + thứ_hạng) theo từng bảng nó xuất hiện → điểm RRF"],
        ["4", "Chọn top-k theo RRF", "Sắp xếp theo điểm RRF, lấy k đoạn đầu (thứ tự này bắt được exact-match)"],
        ["5", "Gán điểm cosine", "Mỗi kết quả giữ điểm cosine similarity (đoạn BM25-only ngoài pool → tính cosine riêng)"],
    ],
)
bold_p("Điểm mấu chốt về gating:")
p("RRF chỉ quyết định THỨ TỰ kết quả. Trường score của mỗi kết quả VẪN là cosine similarity "
  "— nhờ vậy toàn bộ logic chặn kết quả kém (0.3 / 0.5) và tính độ tin cậy (0.8 / 0.5) của "
  "Phase 1 KHÔNG phải sửa gì.")

h2("3.3. Ví dụ minh họa (đã kiểm chứng)")
p("Với câu hỏi \"Điều 650 thừa kế\", trong tình huống vector xếp Điều 650 đứng cuối:")
bullet("Vector-only: [Điều 1, Điều 9, Điều 8, Điều 650] — Điều 650 bị đẩy xuống cuối.")
bullet("Hybrid: [Điều 1, Điều 650, Điều 9, Điều 8] — BM25 khớp token \"650\" đẩy Điều 650 lên hạng 2.")
p("Điểm cosine vẫn được bảo toàn cho từng kết quả, nên gating hoạt động bình thường.")

h2("3.4. Cache chỉ mục BM25")
bullet("Inverted index được dựng lười (lazy) ở lần tìm kiếm hybrid đầu tiên, rồi cache trong bộ nhớ.")
bullet("Khi kho dữ liệu thay đổi (thêm add_chunks, xóa delete_by_so_hieu), cache bị vô hiệu hóa "
       "để lần tìm kiếm kế tiếp dựng lại từ dữ liệu mới.")
bullet("Với quy mô POC (vài văn bản, vài trăm đoạn), chi phí dựng lại không đáng kể.")

h2("3.5. Quyết định thiết kế")
add_table(
    ["Quyết định", "Lựa chọn", "Lý do"],
    [
        ["Thuật toán từ khóa", "BM25 Okapi (k1=1.5, b=0.75)", "Chuẩn công nghiệp; mạnh ở exact-match; không cần huấn luyện"],
        ["Cách trộn xếp hạng", "RRF (k=60)", "Gộp theo thứ hạng, không phụ thuộc thang điểm; đơn giản, ổn định"],
        ["Tokenize", "Bỏ dấu + giữ số", "Câu có/không dấu cùng token; giữ số để bắt \"650\", \"60\""],
        ["Điểm trả về", "Giữ cosine similarity", "Không phá gating & độ tin cậy của Phase 1"],
        ["Pool ứng viên", "max(k×4, 40)", "Đủ rộng để hai nguồn có phần giao nhau cho RRF"],
        ["Cache BM25", "Dựng lười + vô hiệu khi add/delete", "Tránh dựng lại mỗi truy vấn; luôn nhất quán với dữ liệu"],
        ["Cài đặt", "Python thuần (không thư viện mới)", "Nhẹ, không thêm phụ thuộc, dễ bảo trì"],
    ],
)

# ============================================================
# 4. THAY ĐỔI MÃ NGUỒN
# ============================================================
h1("4. Thay đổi mã nguồn")
p("Phase 2 chỉ chạm 2 file, đúng như phạm vi thiết kế.")

h2("4.1. backend/app/storage/vector_store.py")
add_table(
    ["Thành phần thêm/sửa", "Mô tả"],
    [
        ["_remove_diacritics()", "Bỏ dấu tiếng Việt (NFD + xử lý đ/Đ), giữ chữ và số"],
        ["_tokenize()", "Lowercase → bỏ dấu → tách token chữ/số (giữ con số)"],
        ["class _BM25", "Inverted index + Okapi BM25 (k1=1.5, b=0.75); hàm search() trả top-k"],
        ["_ensure_bm25() / _invalidate_bm25()", "Dựng lười + vô hiệu hóa cache chỉ mục BM25"],
        ["Tách search()", "Thành _vector_search() (như cũ) và _hybrid_search() (vector + BM25 + RRF)"],
        ["_cosine_for_ids() / _cosine()", "Tính cosine cho đoạn BM25-only nằm ngoài pool vector"],
        ["Hook add_chunks / delete_by_so_hieu", "Gọi _invalidate_bm25() sau khi corpus đổi"],
    ],
)

h2("4.2. backend/app/pipelines/query_pipeline.py")
add_table(
    ["Thay đổi", "Mô tả"],
    [
        ["Truyền query_text", "Gọi search(query_vector, k=..., query_text=cau_hoi) để kích hoạt nhánh hybrid"],
    ],
)
bold_p("Không đổi: frontend, API routes, schema, system prompt, config (.env).")

# ============================================================
# 5. KIỂM CHỨNG
# ============================================================
h1("5. Kiểm chứng")
add_table(
    ["Hạng mục", "Cách kiểm", "Kết quả"],
    [
        ["Tokenize giữ số + bỏ dấu", "\"Điều 650 thừa kế\" → token", "[\"dieu\",\"650\",\"thua\",\"ke\"] đúng"],
        ["BM25 exact-match", "Corpus 3 đoạn, truy vấn \"Điều 650 thừa kế\"", "Đoạn chứa \"650\" xếp hạng 1"],
        ["Hybrid cứu exact-match", "Vector xếp Điều 650 cuối → chạy hybrid", "Điều 650 lên hạng 2 (cao hơn vector-only)"],
        ["Bảo toàn cosine cho gating", "Kiểm tra score mỗi kết quả", "Nằm trong [-1, 1], đúng cosine"],
        ["Vô hiệu cache khi xóa", "delete_by_so_hieu sau khi đã dựng index", "Cache BM25 bị clear, dựng lại lần sau"],
        ["Import 2 module", "Nạp vector_store, query_pipeline", "OK, không lỗi"],
    ],
)

# ============================================================
# 6. CÔNG NGHỆ SỬ DỤNG
# ============================================================
h1("6. Công nghệ sử dụng")
p("Phase 2 KHÔNG thêm thư viện hay dịch vụ mới. BM25, tokenize và RRF đều viết bằng thư "
  "viện chuẩn của Python.")
add_table(
    ["Thành phần", "Công nghệ", "Ghi chú"],
    [
        ["BM25 + inverted index", "Python thuần (math)", "Okapi BM25, k1=1.5, b=0.75"],
        ["Bỏ dấu tiếng Việt", "unicodedata (chuẩn hóa NFD)", "Xử lý riêng đ/Đ"],
        ["Tokenize", "re (regex)", "Tách token chữ/số"],
        ["RRF merge", "Python thuần", "Hằng số k=60"],
        ["Vector search", "ChromaDB (giữ nguyên Phase 1)", "Cosine similarity"],
    ],
)

# ============================================================
# 7. KẾT LUẬN
# ============================================================
h1("7. Kết luận tính khả thi Phase 2")
add_table(
    ["Khía cạnh", "Đánh giá", "Ghi chú"],
    [
        ["Kỹ thuật", "Khả thi cao", "BM25 + RRF là kỹ thuật chuẩn, cài đặt gọn bằng Python thuần"],
        ["Phạm vi", "Nhỏ, cô lập", "Chỉ 2 file; không đụng API/frontend/prompt/config"],
        ["Rủi ro", "Thấp", "Giữ nguyên gating; có nhánh vector-only dự phòng"],
        ["Chất lượng", "Tăng rõ", "Bắt được exact-match (số Điều, con số) mà Phase 1 miss"],
        ["Mở rộng", "Dễ", "Nền tảng tốt cho Phase 4 (verify trên kết quả hybrid)"],
    ],
)

h2("7.1. Hạn chế còn lại (chuyển sang các Phase sau)")
add_table(
    ["Hạn chế", "Hướng xử lý"],
    [
        ["Câu gõ không dấu vào embedding vẫn kém (BM25 đã bỏ dấu nhưng vector thì chưa)", "Phase 3 — normalize câu không dấu bằng LLM"],
        ["Chưa xác minh chéo câu trả lời của LLM", "Phase 4 — cross-verify bằng LLM thứ 2"],
        ["Chunk chưa được làm giàu (enrich) để tăng recall câu hỏi đời thường", "Phase 4 — enrich chunk khi seed"],
        ["Chưa có xác thực cho upload/xóa", "Phase 5 — auth + admin"],
    ],
)

# Footer
doc.add_paragraph()
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = para.add_run("Ngày lập: Tháng 8/2026")
run.italic = True

# --- Save ---
out = Path(__file__).parent / "bao-cao-phase2-poc.docx"
doc.save(str(out))
print(f"OK: {out}")
